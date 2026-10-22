import argparse
import math
import os
import re
import time
from datetime import datetime
from zoneinfo import ZoneInfo

import serial
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

IST = ZoneInfo("Asia/Kolkata")
HEADERS = [
    "Timestamp (IST)", "Unix Timestamp (ns)", "Serial Arrival Time (monotonic s)",
    "ESP32 Uptime (ms)", "Voltage (V)", "Current (A)", "Active Power (W)",
    "Power Factor", "Energy (Wh)", "Frequency (Hz)", "Apparent Power (VA)",
    "Reactive Power (VAr)", "Estimated Phase Angle (degrees)", "Phase Angle Status",
    "Temperature (°C)", "Serial Framing / Recovery Note",
]
WIDTHS = [27, 23, 29, 20, 16, 16, 20, 16, 15, 17, 23, 22, 34, 43, 19, 42]
EXPECTED_FIELDS = 11


def open_excel_file(filename):
    if os.path.exists(filename):
        workbook = load_workbook(filename)
        if "Measurements" not in workbook.sheetnames:
            worksheet = workbook.create_sheet("Measurements")
            worksheet.append(HEADERS)
        else:
            worksheet = workbook["Measurements"]
            existing = [worksheet.cell(row=1, column=i).value for i in range(1, len(HEADERS) + 1)]
            if existing != HEADERS:
                workbook.close()
                raise ValueError("Existing workbook has a different schema. Use a NEW output filename, e.g. pzem_serial_diagnostics.xlsx.")
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Measurements"
        worksheet.append(HEADERS)

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, width in enumerate(WIDTHS, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.row_dimensions[1].height = 34
    return workbook, worksheet


def parse_candidate(candidate):
    parts = [part.strip() for part in candidate.strip().split(",")]
    if len(parts) != EXPECTED_FIELDS or parts[0] != "DATA":
        return None, f"expected {EXPECTED_FIELDS} fields, got {len(parts)}"
    try:
        values = [float(value) for value in parts[1:]]
    except ValueError as exc:
        return None, f"non-numeric field: {exc}"
    (voltage, current, active_power, power_factor, energy, frequency,
     apparent_power, reactive_power, esp32_millis, temperature_c) = values

    def finite_or_none(value):
        return value if math.isfinite(value) else None

    try:
        uptime = int(esp32_millis)
    except (ValueError, OverflowError):
        return None, "invalid ESP32 uptime"

    return {
        "voltage": finite_or_none(voltage),
        "current": finite_or_none(current),
        "active_power": finite_or_none(active_power),
        "power_factor": finite_or_none(power_factor),
        "energy": finite_or_none(energy),
        "frequency": finite_or_none(frequency),
        "apparent_power": finite_or_none(apparent_power),
        "reactive_power": finite_or_none(reactive_power),
        "esp32_millis": uptime,
        "temperature_c": finite_or_none(temperature_c),
    }, None


def extract_records(text):
    valid, notes, non_data = [], [], []
    positions = [m.start() for m in re.finditer(r"(?<![A-Za-z0-9_])DATA,", text)]
    if not positions:
        if text.strip():
            non_data.append(text.strip())
        return valid, notes, non_data
    if text[:positions[0]].strip():
        non_data.append(text[:positions[0]].strip())

    candidates = [
        text[start:(positions[i + 1] if i + 1 < len(positions) else len(text))].strip(" \r\n\t")
        for i, start in enumerate(positions)
    ]
    for candidate in candidates:
        record, error = parse_candidate(candidate)
        if record is None:
            notes.append(f"Rejected candidate: {error}: {candidate!r}")
        else:
            valid.append((record, "Recovered by splitting concatenated DATA records" if len(candidates) > 1 else "OK"))
    return valid, notes, non_data


def main():
    parser = argparse.ArgumentParser(description="VoltGuard PZEM + DS18B20 serial diagnostics and Excel logger")
    parser.add_argument("--port", required=True, help="ESP32 COM port, e.g. COM17")
    parser.add_argument("--baud", type=int, default=115200)
    parser.add_argument("--out", default="pzem_serial_diagnostics.xlsx", help="Use a NEW filename for this schema")
    parser.add_argument("--timeout", type=float, default=2.0)
    args = parser.parse_args()
    filename = os.path.abspath(args.out)

    print("=" * 70)
    print("VOLTGUARD: SERIAL DIAGNOSTICS + PZEM/DS18B20 EXCEL LOGGER")
    print("=" * 70)
    print(f"Serial port: {args.port}")
    print(f"Baud rate:   {args.baud}")
    print("Time zone:   Asia/Kolkata")
    print(f"Excel file:  {filename}")
    print("Expected DATA fields: 11 total (DATA + 10 numeric fields)")
    print("Records are timestamped on arrival; no artificial samples are created.")
    print("Press Ctrl+C to stop.")
    print("=" * 70)

    workbook = None
    count = 0
    rejected_count = 0
    last_arrival_mono = None
    last_uptime = None
    text_buffer = ""

    try:
        workbook, worksheet = open_excel_file(filename)
        with serial.Serial(args.port, args.baud, timeout=args.timeout) as esp32:
            time.sleep(0.2)
            print("\nConnected to ESP32.")
            print("Inspecting serial line endings and records...\n")

            while True:
                first = esp32.read(1)
                if not first:
                    continue
                waiting = esp32.in_waiting
                chunk = first + (esp32.read(waiting) if waiting else b"")
                arrival_mono = time.monotonic()
                arrival_ist = datetime.now(IST).replace(tzinfo=None)
                arrival_unix_ns = time.time_ns()
                decoded = chunk.decode("utf-8", errors="replace")
                text_buffer += decoded

                while "\n" in text_buffer:
                    raw_segment, text_buffer = text_buffer.split("\n", 1)
                    segment = raw_segment.rstrip("\r")
                    if not segment.strip():
                        continue
                    if segment.startswith("#"):
                        print(segment)
                        continue

                    valid, notes, non_data = extract_records(segment)
                    for message in non_data:
                        print(f"SERIAL TEXT (not DATA): {message!r}")
                    for message in notes:
                        rejected_count += 1
                        print(f"WARNING: {message}")

                    for record, recovery_note in valid:
                        pf = record["power_factor"]
                        if pf is not None and -1.0 <= pf <= 1.0:
                            theta = math.degrees(math.acos(max(-1.0, min(1.0, pf))))
                            phase_status = "PF-derived estimate; not measured phasor"
                        else:
                            theta = None
                            phase_status = "Unavailable"

                        gap = None if last_arrival_mono is None else arrival_mono - last_arrival_mono
                        uptime_gap = None if last_uptime is None else record["esp32_millis"] - last_uptime
                        note = recovery_note
                        if gap is not None and gap > 2.5:
                            note += f"; arrival gap {gap:.3f}s"
                        if uptime_gap is not None and uptime_gap < 0:
                            note += "; ESP32 uptime reset/wrap"
                        elif uptime_gap is not None and uptime_gap > 2500:
                            note += f"; ESP32 uptime gap {uptime_gap}ms"

                        worksheet.append([
                            arrival_ist, arrival_unix_ns, round(arrival_mono, 6),
                            record["esp32_millis"], record["voltage"], record["current"],
                            record["active_power"], pf, record["energy"], record["frequency"],
                            record["apparent_power"], record["reactive_power"], theta,
                            phase_status, record["temperature_c"], note,
                        ])
                        row = worksheet.max_row
                        worksheet.cell(row, 1).number_format = "yyyy-mm-dd hh:mm:ss.000"
                        worksheet.cell(row, 2).number_format = "0"
                        worksheet.cell(row, 3).number_format = "0.000000"
                        worksheet.cell(row, 4).number_format = "0"
                        for col in (5, 7, 13, 15):
                            worksheet.cell(row, col).number_format = "0.00"
                        for col in (6, 8):
                            worksheet.cell(row, col).number_format = "0.0000"

                        workbook.save(filename)
                        count += 1
                        gap_text = "first record" if gap is None else f"{gap:.3f}s since previous accepted record"
                        print(f"Record {count} | {arrival_ist.strftime('%Y-%m-%d %H:%M:%S.%f')[:-3]} IST | uptime {record['esp32_millis']} ms | {gap_text}")
                        print(f"  V={record['voltage']} V, I={record['current']} A, P={record['active_power']} W, PF={pf}, T={record['temperature_c']} °C")
                        if recovery_note != "OK":
                            print(f"  NOTE: {recovery_note}")
                        print(f"  Excel row: {row}")
                        last_arrival_mono = arrival_mono
                        last_uptime = record["esp32_millis"]

    except KeyboardInterrupt:
        print("\nLogging stopped by user.")
    except serial.SerialException as error:
        print("\nSerial port error:", error)
        print("Close Arduino Serial Monitor/Plotter and check the COM port.")
    except PermissionError:
        print("\nCannot access the Excel file. Close it in Excel or choose another output filename.")
    except ValueError as error:
        print("\nWorkbook/configuration error:", error)
    finally:
        if workbook is not None:
            try:
                workbook.save(filename)
                workbook.close()
                print(f"\nExcel file saved: {filename}")
                print(f"Valid records captured this session: {count}")
                print(f"Rejected malformed candidates: {rejected_count}")
                if text_buffer.strip():
                    print(f"Trailing incomplete serial buffer at stop: {text_buffer!r}")
            except PermissionError:
                print("\nCould not save workbook. Close it in Excel and check permissions.")


if __name__ == "__main__":
    main()
